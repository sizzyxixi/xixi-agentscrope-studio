"""Format-aware processing; no arbitrary shell/Python execution is exposed to agents."""

from __future__ import annotations
import asyncio
import csv
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from xml.sax.saxutils import escape

TEXT = set(
    "txt log rst md markdown html htm js ts py sql json jsonl ndjson yaml yml xml css jsx tsx sh c cpp h java go rs toml ini env conf cfg svg".split()
)
TEXT |= set(
    "mjs cjs mts cts pyi pyw ipynb scss sass less bash zsh fish ksh cc cxx hpp hh cs php rb swift kt kts dart vue svelte astro scala lua r jl pl pm ps1 bat cmd gradle lock diff patch cmake tf hcl proto graphql gql properties".split()
)
IMAGES = {"png", "jpg", "jpeg", "webp", "gif", "bmp", "tif", "tiff", "avif"}
AUDIO = set("mp3 wav m4a aac flac ogg opus weba".split())
VIDEO = set("mp4 mov webm mkv m4v ogv avi mpeg mpg".split())
LEGACY = {
    "doc": "docx",
    "odt": "docx",
    "rtf": "docx",
    "ppt": "pptx",
    "odp": "pptx",
    "xls": "xlsx",
    "ods": "xlsx",
}
MAX_PARSE = 32_000_000
MAX_TEXT = 240_000
MAX_ROWS, MAX_COLS = 500, 100


def ext(name):
    return Path(name).suffix.lower().lstrip(".")


def decode(data: bytes):
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    return data.decode("utf-8", errors="replace")


def safe_archive(path: Path, max_total=256_000_000):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(i.file_size for i in entries) > max_total:
            raise ValueError("压缩内容超出当前处理范围。")
        for item in entries:
            name = item.filename.replace("\\", "/")
            if (
                name.startswith("/")
                or ".." in PurePosixPath(name).parts
                or len(PurePosixPath(name).parts) > 10
                or item.flag_bits & 1
            ):
                raise ValueError("不支持加密压缩包或非法目录。")
            if item.external_attr >> 16 & 0o170000 == 0o120000:
                raise ValueError("不支持符号链接。")
        return [
            {"name": i.filename, "size": i.file_size, "directory": i.is_dir()}
            for i in entries
        ]


def office_convert(source: Path, name: str, target: str, destination: Path):
    binary = (
        os.getenv("AGENTSCOPE_SOFFICE")
        or shutil.which("soffice")
        or shutil.which("libreoffice")
    )
    if not binary:
        raise ValueError("当前未部署 Office 格式转换引擎，无法完成此转换。")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        incoming = root / ("input." + ext(name))
        shutil.copyfile(source, incoming)
        out = root / "out"
        out.mkdir()
        command = [
            binary,
            f'-env:UserInstallation={root.joinpath("profile").as_uri()}',
            "--headless",
            "--convert-to",
            target,
            "--outdir",
            str(out),
            str(incoming),
        ]
        try:
            result = subprocess.run(command, capture_output=True, timeout=90)
        except (OSError, subprocess.TimeoutExpired) as e:
            raise ValueError("Office 文件转换失败，原文件已保留。") from e
        outputs = list(out.glob("*." + target))
        if result.returncode or not outputs:
            raise ValueError("Office 文件转换失败，原文件已保留。")
        shutil.copyfile(outputs[0], destination)


def ffmpeg():
    import imageio_ffmpeg

    return (
        os.getenv("AGENTSCOPE_FFMPEG")
        or shutil.which("ffmpeg")
        or imageio_ffmpeg.get_ffmpeg_exe()
    )


def media_convert(source: Path, target: str, output: Path):
    args = (
        ["-vn", "-c:a", "pcm_s16le"]
        if target == "wav"
        else (
            ["-vn", "-c:a", "libmp3lame"]
            if target == "mp3"
            else (
                ["-c:v", "libvpx-vp9", "-c:a", "libopus"]
                if target == "webm"
                else ["-c:v", "libx264", "-c:a", "aac", "-movflags", "+faststart"]
            )
        )
    )
    try:
        result = subprocess.run(
            [ffmpeg(), "-nostdin", "-y", "-i", str(source), *args, str(output)],
            capture_output=True,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        raise ValueError("音视频转换失败，原文件已保留。") from e
    if result.returncode or not output.exists():
        raise ValueError("音视频编码不支持或文件损坏，转换失败。")


def svg_png(data: bytes):
    from defusedxml import ElementTree

    root = ElementTree.fromstring(data)
    if root.tag.split("}")[-1] != "svg":
        raise ValueError("无效 SVG。")
    for element in root.iter():
        if element.tag.split("}")[-1] in {
            "script",
            "foreignObject",
            "image",
            "use",
            "style",
        }:
            raise ValueError("SVG 含不支持的活动内容或资源引用。")
        for key, value in element.attrib.items():
            if (
                key.lower().startswith("on")
                or "href" in key.lower()
                or "url(" in value.lower()
                or "@import" in value.lower()
            ):
                raise ValueError("SVG 含外部资源或脚本。")
    import resvg_py

    return resvg_py.svg_to_bytes(
        svg_string=data.decode("utf-8"), width=1200, height=1200, skip_system_fonts=True
    )


def ocr(image):
    url = os.getenv("AGENTSCOPE_OCR_URL", "").rstrip("/")
    if not url:
        from .ocr_service.engine import recognize

        return recognize(image)
    import hashlib
    import hmac
    import httpx
    from PIL import Image, ImageOps

    key = os.environ.get("AGENTSCOPE_FILE_KEY", "")
    if not key:
        raise ValueError("OCR 服务缺少身份验证配置。")
    picture = ImageOps.exif_transpose(image).convert("RGBA")
    picture.thumbnail((2400, 2400), Image.Resampling.LANCZOS)
    background = Image.new("RGB", picture.size, "white")
    background.paste(picture, mask=picture.getchannel("A"))
    payload = io.BytesIO()
    background.save(payload, format="JPEG", quality=95)
    data = payload.getvalue()
    signature = hmac.new(key.encode(), b"ocr-v1:" + data, hashlib.sha256).hexdigest()
    try:
        response = httpx.post(
            url + "/recognize", content=data,
            headers={"Content-Type": "image/jpeg", "X-OCR-Signature": signature},
            timeout=120,
        )
        response.raise_for_status()
        return response.json()["text"]
    except (httpx.HTTPError, KeyError, ValueError) as e:
        raise ValueError("OCR 识别服务暂不可用，请稍后重试。") from e


def transcribe(path: Path):
    from faster_whisper import WhisperModel

    model = WhisperModel(
        os.getenv("AGENTSCOPE_WHISPER_MODEL", "tiny"), device="cpu", compute_type="int8"
    )
    segments, _ = model.transcribe(str(path), beam_size=3)
    text = []
    for segment in segments:
        text.append(f"[{segment.start:.1f}s] {segment.text}")
        if sum(map(len, text)) > MAX_TEXT:
            break
    return "\n".join(text)


async def extract(path: Path, name: str, *, offset=0, limit=12000):
    suffix = ext(name)
    note = ""
    text = ""
    if path.stat().st_size > MAX_PARSE and suffix not in AUDIO | VIDEO:
        raise ValueError("文件超出当前解析范围（32 MB）；原文件仍可下载。")
    if suffix in {"pdf", "docx", "pptx", "xlsx", "xls"}:
        from agentscope.rag import PDFParser, WordParser, PPTParser, ExcelParser

        if suffix in {"docx", "pptx", "xlsx"}:
            safe_archive(path)
        parser = {
            "pdf": PDFParser,
            "docx": WordParser,
            "pptx": PPTParser,
            "xlsx": ExcelParser,
            "xls": ExcelParser,
        }[suffix]()
        sections = await parser.parse(str(path), filename=name)
        texts = []
        image_count = 0
        for section in sections:
            content = section.content
            for block in content if isinstance(content, list) else [content]:
                if getattr(block, "text", None):
                    texts.append(block.text)
                elif (
                    getattr(
                        getattr(block, "source", None), "media_type", ""
                    ).startswith("image/")
                    and getattr(block.source, "type", "") == "base64"
                ):
                    if image_count >= 20:
                        note += " 仅提取前 20 张内嵌图片的文字。"
                        continue
                    image_count += 1
                    try:
                        import base64
                        from PIL import Image

                        picture = Image.open(
                            io.BytesIO(base64.b64decode(block.source.data))
                        )
                        recognized = await asyncio.to_thread(ocr, picture)
                        if recognized:
                            texts.append("[内嵌图片 OCR，可能存在误差]\n" + recognized)
                    except Exception:
                        note += " 有内嵌图片未能提取文字。"
        text = "\n\n".join(texts)
        if suffix == "xlsx" and not text.strip():
            text = json.dumps(table_preview(path, name), ensure_ascii=False)
        if suffix == "pdf" and any(
            not getattr(section.content, "text", "").strip() for section in sections
        ):
            import pymupdf

            from PIL import Image

            page_text = [getattr(section.content, "text", "") for section in sections]
            with pymupdf.open(path) as document:
                scanned = [
                    index
                    for index, value in enumerate(page_text)
                    if not value.strip() and document[index].get_images()
                ]
                for index in scanned[:20]:
                    page = document[index]
                    scale = min(1.5, 2400 / max(page.rect.width, page.rect.height))
                    with Image.open(io.BytesIO(
                        page.get_pixmap(matrix=pymupdf.Matrix(scale, scale)).tobytes("png")
                    )) as picture:
                        page_text[index] = await asyncio.to_thread(ocr, picture)
            text = "\n\n".join(page_text)
            if scanned:
                note = "OCR 结果可能存在识别误差；最多处理 20 个扫描页面。"
    elif suffix in LEGACY:
        with tempfile.TemporaryDirectory() as tmp:
            converted = Path(tmp) / ("converted." + LEGACY[suffix])
            await asyncio.to_thread(
                office_convert, path, name, LEGACY[suffix], converted
            )
            return await extract(converted, converted.name, offset=offset, limit=limit)
    elif suffix in IMAGES:
        from PIL import Image

        with Image.open(path) as picture:
            text = await asyncio.to_thread(ocr, picture)
        note = "图片文字由 OCR 提取，识别可能存在误差；不代表对图片含义的视觉理解。"
    elif suffix in AUDIO | VIDEO:
        text = await asyncio.to_thread(transcribe, path)
        note = "自动转录可能存在识别误差。"
        if suffix in VIDEO:
            note += " 视频附带最多 6 个采样画面的 OCR 文本。"
            with tempfile.TemporaryDirectory() as tmp:
                try:
                    await asyncio.to_thread(
                        subprocess.run,
                        [
                            ffmpeg(),
                            "-nostdin",
                            "-y",
                            "-i",
                            str(path),
                            "-vf",
                            "fps=1/30",
                            "-frames:v",
                            "6",
                            str(Path(tmp) / "%02d.png"),
                        ],
                        capture_output=True,
                        timeout=120,
                    )
                    from PIL import Image

                    for frame in sorted(Path(tmp).glob("*.png")):
                        text += "\n画面文字：" + await asyncio.to_thread(
                            ocr, Image.open(frame)
                        )
                except Exception:
                    note += " 画面提取不可用。"
    elif suffix in {"zip", "epub"}:
        entries = safe_archive(path)
        text = json.dumps(entries, ensure_ascii=False)
        if suffix == "epub":
            from bs4 import BeautifulSoup

            with zipfile.ZipFile(path) as archive:
                text = "\n".join(
                    BeautifulSoup(archive.read(e["name"]), "html.parser").get_text(
                        " ", strip=True
                    )
                    for e in entries
                    if e["name"].endswith((".xhtml", ".html", ".htm"))
                )
    elif (
        suffix in TEXT
        or suffix in {"csv", "tsv", "xlsm", "ods"}
        or name in {"Dockerfile", "Makefile", ".env"}
    ):
        if suffix in {"xlsm", "ods"}:
            data = table_preview(path, name)
            text = json.dumps(data, ensure_ascii=False)
        elif suffix in {"html", "htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(decode(path.read_bytes()), "html.parser")
            for element in soup(["script", "style"]):
                element.decompose()
            text = soup.get_text("\n", strip=True)
        else:
            text = decode(path.read_bytes())
    else:
        raise ValueError("当前没有此文件格式的内容提取能力；原文件仍可下载。")
    end = min(offset + min(max(limit, 1), 24000), len(text), MAX_TEXT)
    return {
        "text": text[offset:end],
        "offset": offset,
        "next_offset": end if end < min(len(text), MAX_TEXT) else None,
        "total_characters": len(text),
        "notice": note + (" 仅展示部分内容。" if end < len(text) else ""),
    }


def table_preview(path: Path, name: str):
    suffix = ext(name)
    sheets = []
    truncated = False
    if suffix in {"xlsx", "xlsm"}:
        import openpyxl

        safe_archive(path)
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for sheet in values.worksheets[:30]:
            if sheet.sheet_state != "visible":
                continue
            rows = []
            max_row = min(sheet.max_row or 0, MAX_ROWS)
            max_col = min(sheet.max_column or 0, MAX_COLS)
            formula_rows = formulas[sheet.title].iter_rows(
                min_row=1, max_row=max_row, max_col=max_col
            )
            for row, formula_row in zip(
                sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col),
                formula_rows,
            ):
                cells = []
                for cell, formula in zip(row, formula_row):
                    value = cell.value
                    if value is None and formula.data_type == "f":
                        value = f"{formula.value}（未计算）"
                    cells.append("" if value is None else str(value))
                rows.append(cells)
            sheets.append(
                {"name": sheet.title, "rows": rows, "total_rows": sheet.max_row}
            )
            truncated |= (sheet.max_row or 0) > MAX_ROWS or (
                sheet.max_column or 0
            ) > MAX_COLS
        values.close()
        formulas.close()
    elif suffix == "ods":
        from odf.opendocument import load
        from odf.table import Table, TableRow, TableCell
        from odf import teletype

        safe_archive(path)
        for sheet in load(str(path)).spreadsheet.getElementsByType(Table)[:30]:
            rows = []
            for row in sheet.getElementsByType(TableRow)[:MAX_ROWS]:
                cells = []
                for cell in row.getElementsByType(TableCell)[:MAX_COLS]:
                    repeat = min(
                        int(cell.getAttribute("numbercolumnsrepeated") or 1),
                        MAX_COLS - len(cells),
                    )
                    cells.extend([teletype.extractText(cell)] * repeat)
                rows.append(cells[:MAX_COLS])
            sheets.append(
                {
                    "name": sheet.getAttribute("name"),
                    "rows": rows,
                    "total_rows": len(rows),
                }
            )
        truncated = True
    else:
        delimiter = "\t" if suffix == "tsv" else ","
        reader = csv.reader(io.StringIO(decode(path.read_bytes())), delimiter=delimiter)
        rows = []
        for index, row in enumerate(reader):
            if index >= MAX_ROWS:
                truncated = True
                break
            truncated |= len(row) > MAX_COLS
            rows.append(row[:MAX_COLS])
        sheets = [{"name": Path(name).stem, "rows": rows, "total_rows": None}]
    return {
        "kind": "table",
        "sheets": sheets,
        "notice": "仅展示部分内容" if truncated else "",
    }


def structured_content(content):
    try:
        return json.loads(content)
    except (ValueError, TypeError) as e:
        raise ValueError("此格式需要 JSON 结构化内容。") from e


def render(name: str, content: str, output: Path, assets: dict | None = None):
    suffix = ext(name)
    assets = assets or {}
    if suffix in {"docx", "pdf"} and len(content.splitlines()) > 10000:
        raise ValueError("单次文档生成最多 10000 段，请分批处理。")
    if suffix in TEXT:
        if suffix == "json":
            json.loads(content)
        elif suffix in {"jsonl", "ndjson"}:
            for line in content.splitlines():
                if line.strip():
                    json.loads(line)
        elif suffix in {"yaml", "yml"}:
            import yaml

            yaml.safe_load(content)
        elif suffix == "xml":
            from defusedxml import ElementTree

            ElementTree.fromstring(content)
        elif suffix == "svg":
            svg_png(content.encode())
        output.write_text(content, encoding="utf-8")
    elif suffix in {"csv", "tsv"}:
        rows = (
            structured_content(content)
            if content.lstrip().startswith("[")
            else list(
                csv.reader(
                    io.StringIO(content), delimiter="\t" if suffix == "tsv" else ","
                )
            )
        )
        if not isinstance(rows, list):
            raise ValueError("表格内容需要二维数组。")
        with output.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, delimiter="\t" if suffix == "tsv" else ",").writerows(rows)
    elif suffix == "docx":
        from docx import Document
        from docx.shared import Inches
        from docx.oxml.ns import qn

        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Arial"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        structured = (
            structured_content(content) if content.lstrip().startswith("{") else None
        )
        body = (
            "\n".join([structured.get("title", ""), *structured.get("paragraphs", [])])
            if isinstance(structured, dict)
            else content
        )
        for line in body.splitlines():
            match = re.match(r"^(#{1,6}) (.*)", line)
            if match:
                document.add_heading(match[2], len(match[1]))
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("|") and line.endswith("|"):
                cells = [x.strip() for x in line.strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?", x) for x in cells):
                    continue
                table = document.add_table(rows=1, cols=len(cells))
                table.style = "Table Grid"
                for cell, value in zip(table.rows[0].cells, cells):
                    cell.text = value
            else:
                document.add_paragraph(line)
        if isinstance(structured, dict):
            for rows in structured.get("tables", []):
                if not rows:
                    continue
                table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = str(value)
            for image in structured.get("images", []):
                identity = image.get("file_id")
                if identity not in assets:
                    raise ValueError("图片资源不可用。")
                document.add_picture(
                    str(assets[identity]),
                    width=Inches(min(float(image.get("width_inches", 5)), 7)),
                )
                if image.get("caption"):
                    document.add_paragraph(image["caption"])
        document.save(output)
    elif suffix == "pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle(
            "body", fontName="STSong-Light", fontSize=11, leading=18, wordWrap="CJK"
        )
        flow = []
        for line in content.splitlines():
            flow.extend([Paragraph(escape(line) or " ", style), Spacer(1, 5)])
        SimpleDocTemplate(str(output)).build(flow or [Paragraph(" ", style)])
    elif suffix == "xlsx":
        import openpyxl

        data = structured_content(content)
        if isinstance(data, list):
            data = {"Sheet1": data}
        if not isinstance(data, dict) or not data or len(data) > 30:
            raise ValueError("Excel 内容为工作表名称到二维数组的 JSON 对象。")
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, rows in data.items():
            settings = rows if isinstance(rows, dict) else {}
            rows = settings.get("rows", []) if settings else rows
            sheet = book.create_sheet(str(name)[:31])
            if len(rows) > 10000:
                raise ValueError("每张工作表最多生成 10000 行。")
            for row in rows:
                sheet.append(row)
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 20
            sheet.freeze_panes = "A2"
            for image in settings.get("images", []):
                from openpyxl.drawing.image import Image

                identity = image.get("file_id")
                if identity not in assets:
                    raise ValueError("图片资源不可用。")
                picture = Image(str(assets[identity]))
                sheet.add_image(picture, image.get("anchor", "A1"))
        book.save(output)
    elif suffix == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        data = structured_content(content)
        slides = data.get("slides", []) if isinstance(data, dict) else data
        if not isinstance(slides, list) or not slides or len(slides) > 100:
            raise ValueError(
                'PPT 内容为 {"slides":[{"title":"标题","body":"正文"}]}，最多 100 页。'
            )
        deck = Presentation()
        for item in slides:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = item.get("title", "")
            slide.placeholders[1].text = item.get("body", "")
            for image in item.get("images", []):
                identity = image.get("file_id")
                if identity not in assets:
                    raise ValueError("图片资源不可用。")
                slide.shapes.add_picture(
                    str(assets[identity]),
                    Inches(float(image.get("left", 1))),
                    Inches(float(image.get("top", 2))),
                    width=Inches(float(image.get("width", 4))),
                )
            for chart in item.get("charts", []):
                from pptx.chart.data import CategoryChartData
                from pptx.enum.chart import XL_CHART_TYPE

                values = CategoryChartData()
                values.categories = chart.get("categories", [])
                for series in chart.get("series", []):
                    values.add_series(
                        series.get("name", "数据"), series.get("values", [])
                    )
                chart_type = {
                    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                    "line": XL_CHART_TYPE.LINE,
                    "pie": XL_CHART_TYPE.PIE,
                }.get(chart.get("type"), XL_CHART_TYPE.COLUMN_CLUSTERED)
                slide.shapes.add_chart(
                    chart_type, Inches(1), Inches(2), Inches(8), Inches(4), values
                )
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.name = "Microsoft YaHei"
        deck.save(output)
    elif suffix in {"png", "jpg", "jpeg", "webp", "gif"}:
        from PIL import Image, ImageDraw, ImageFont

        data = structured_content(content)
        frames = data.get("frames", [data])
        if not frames or len(frames) > 60:
            raise ValueError("GIF 需要 1–60 帧。")
        dimensions = [
            (int(frame.get("width", 1024)), int(frame.get("height", 768)))
            for frame in frames
        ]
        if any(min(size) < 1 or max(size) > 4096 for size in dimensions):
            raise ValueError("图片尺寸须在 1–4096 像素内。")
        if sum(width * height for width, height in dimensions) > 64_000_000:
            raise ValueError("动画总像素超出 6400 万处理范围。")
        images = []
        for frame in frames:
            width, height = int(frame.get("width", 1024)), int(frame.get("height", 768))
            if min(width, height) < 1 or max(width, height) > 4096:
                raise ValueError("图片尺寸须在 1–4096 像素内。")
            image = Image.new("RGB", (width, height), frame.get("background", "white"))
            for source in frame.get("images", []):
                identity = source.get("file_id")
                if identity not in assets:
                    raise ValueError("图片资源不可用。")
                picture = Image.open(assets[identity]).convert("RGBA")
                picture.thumbnail(
                    (
                        min(4096, int(source.get("width", width))),
                        min(4096, int(source.get("height", height))),
                    )
                )
                image.paste(
                    picture,
                    (int(source.get("x", 0)), int(source.get("y", 0))),
                    picture.getchannel("A"),
                )
            draw = ImageDraw.Draw(image)
            for item in frame.get("shapes", [])[:500]:
                box = item.get("box", [20, 20, 200, 100])
                fill = item.get("fill", "#490cf2")
                if item.get("type") == "ellipse":
                    draw.ellipse(box, fill=fill)
                elif item.get("type") == "line":
                    draw.line(box, fill=fill, width=int(item.get("width", 3)))
                else:
                    draw.rectangle(box, fill=fill)
            for item in frame.get("texts", [])[:200]:
                font_path = os.getenv("AGENTSCOPE_FONT") or next(
                    (
                        p
                        for p in [
                            "/System/Library/Fonts/PingFang.ttc",
                            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                        ]
                        if Path(p).exists()
                    ),
                    None,
                )
                font = (
                    ImageFont.truetype(
                        font_path, max(1, min(256, int(item.get("size", 28))))
                    )
                    if font_path
                    else ImageFont.load_default()
                )
                import textwrap

                x, y = int(item.get("x", 20)), int(item.get("y", 20))
                font_size = max(1, min(256, int(item.get("size", 28))))
                lines = textwrap.wrap(
                    str(item.get("text", ""))[:4000],
                    width=max(1, (width - max(0, x)) // font_size),
                )
                visible = "\n".join(
                    lines[: max(1, (height - max(0, y)) // (font_size + 4))]
                )
                draw.multiline_text(
                    (x, y),
                    visible,
                    fill=item.get("color", "#222"),
                    font=font,
                )
            images.append(image)
        images[0].save(
            output,
            format={"jpg": "JPEG", "jpeg": "JPEG"}.get(suffix, suffix.upper()),
            save_all=suffix == "gif",
            append_images=images[1:] if suffix == "gif" else [],
            duration=max(50, int(data.get("duration", 300))),
            loop=0,
        )
    else:
        raise ValueError(f"当前无法生成 {suffix.upper()} 文件。")


def validate(path: Path, name: str):
    suffix = ext(name)
    if suffix == "docx":
        from docx import Document

        Document(path)
    elif suffix == "xlsx":
        import openpyxl

        openpyxl.load_workbook(path, read_only=True).close()
    elif suffix == "pptx":
        from pptx import Presentation

        Presentation(path)
    elif suffix == "pdf":
        from pypdf import PdfReader

        PdfReader(path)
    elif suffix in IMAGES:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
            actual = {"JPEG": "jpg", "PNG": "png", "WEBP": "webp", "GIF": "gif"}.get(
                image.format
            )
            if actual and actual != ("jpg" if suffix == "jpeg" else suffix):
                raise ValueError("文件内容与目标图片格式不一致。")
    elif suffix == "zip":
        safe_archive(path)
    elif suffix == "svg":
        svg_png(path.read_bytes())
    elif suffix in AUDIO | VIDEO:
        result = subprocess.run(
            [ffmpeg(), "-v", "error", "-i", str(path), "-t", "0.1", "-f", "null", "-"],
            capture_output=True,
            timeout=30,
        )
        if result.returncode:
            raise ValueError("生成的音视频无法解码。")
    elif suffix not in TEXT | {"csv", "tsv"}:
        raise ValueError("不支持此目标格式。")
    if not path.exists():
        raise ValueError("文件生成失败。")


async def convert(source: Path, name: str, target: str, output: Path, options: dict):
    original = ext(name)
    target = target.lower().lstrip(".")
    notice = ""
    if target in {"png", "jpg", "jpeg", "webp", "gif"} and original in IMAGES | {"svg"}:
        from PIL import Image

        image = (
            Image.open(io.BytesIO(svg_png(source.read_bytes())))
            if original == "svg"
            else Image.open(source)
        )
        if original == "gif" and target != "gif":
            image.seek(0)
            notice = "已取 GIF 首帧，动画不保留。"
        if target in {"jpg", "jpeg"}:
            rgba = image.convert("RGBA")
            image = Image.new("RGB", rgba.size, "white")
            image.paste(rgba, mask=rgba.getchannel("A"))
            notice += "透明区域使用白底。"
        image.save(
            output, format={"jpg": "JPEG", "jpeg": "JPEG"}.get(target, target.upper())
        )
    elif (
        target in {"txt", "md"} and original in {"docx", "pdf"} | AUDIO | IMAGES | VIDEO
    ):
        result = await extract(source, name, limit=24000)
        parts = [result["text"]]
        while result["next_offset"] is not None:
            result = await extract(
                source, name, offset=result["next_offset"], limit=24000
            )
            parts.append(result["text"])
        output.write_text("\n".join(parts), encoding="utf-8")
        notice = result["notice"]
    elif target in {"docx", "pdf", "html"} and original in {
        "txt",
        "md",
        "markdown",
        "html",
        "htm",
    }:
        content = decode(source.read_bytes())
        if target == "html":
            from markdown_it import MarkdownIt

            output.write_text(
                '<!doctype html><meta charset="utf-8">'
                + MarkdownIt("commonmark", {"html": False}).render(content),
                encoding="utf-8",
            )
        else:
            if original in {"html", "htm"}:
                from bs4 import BeautifulSoup

                content = BeautifulSoup(content, "html.parser").get_text("\n")
            render("output." + target, content, output)
            notice = "按基础文本排版生成；复杂样式不保留。"
    elif (target == "pdf" and original in {"docx", "pptx"}) or (
        original in LEGACY and target == LEGACY[original]
    ):
        await asyncio.to_thread(office_convert, source, name, target, output)
        notice = "静态转换；动画、交互和复杂排版可能变化。"
    elif original in {"xlsx", "csv", "tsv"} and target in {
        "csv",
        "tsv",
        "xlsx",
        "json",
    }:
        data = table_preview(source, name)
        sheet = (
            next((s for s in data["sheets"] if s["name"] == options.get("sheet")), None)
            if options.get("sheet")
            else next(iter(data["sheets"]), None)
        )
        if not sheet:
            raise ValueError("没有可导出的工作表。")
        if data["notice"]:
            raise ValueError(
                "工作表超出完整转换范围，未导出抽样数据；请选择较小工作表。"
            )
        rows = sheet["rows"]
        if target == "json":
            headers = rows[0] if rows else []
            output.write_text(
                json.dumps(
                    [dict(zip(headers, row)) for row in rows[1:]],
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
        else:
            render("output." + target, json.dumps(rows, ensure_ascii=False), output)
        if any("未计算" in str(c) for row in rows for c in row):
            notice = "包含未计算公式，未编造计算结果。"
    elif original in {"json", "yaml", "yml"} and target in {"json", "yaml", "yml"}:
        import yaml

        data = (
            json.loads(decode(source.read_bytes()))
            if original == "json"
            else yaml.safe_load(decode(source.read_bytes()))
        )
        output.write_text(
            (
                json.dumps(data, ensure_ascii=False, indent=2)
                if target == "json"
                else yaml.safe_dump(data, allow_unicode=True)
            ),
            encoding="utf-8",
        )
    elif (
        original in AUDIO | VIDEO
        and target in {"mp3", "wav", "mp4", "webm"}
        and (original in VIDEO or target in {"mp3", "wav"})
    ):
        await asyncio.to_thread(media_convert, source, target, output)
        notice = "重新编码可能影响画质或音质。"
    else:
        raise ValueError(
            f"当前无法将 {original.upper()} 转换为 {target.upper()}，原文件已保留。"
        )
    validate(output, "output." + target)
    return notice
